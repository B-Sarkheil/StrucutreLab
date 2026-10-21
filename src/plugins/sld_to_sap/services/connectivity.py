"""Read frames and joints from the SAP2000 'Export to Excel' file (docs section 3.2)."""
from __future__ import annotations

import threading
from dataclasses import dataclass
from pathlib import Path

from .issues import Issue

FRAME_SHEET = "Connectivity - Frame"
JOINT_SHEET = "Joint Coordinates"


@dataclass(frozen=True)
class Frame:
    name: str
    i: tuple[float, float, float]
    j: tuple[float, float, float]
    length: float


def _name(value) -> str:
    """SAP names may come as numbers (9, 9.0) or text."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header(rows: list[tuple], required: list[str]) -> tuple[int, dict[str, int]] | None:
    for r, row in enumerate(rows[:10]):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if all(name in cells for name in required):
            return r, {name: cells.index(name) for name in required}
    return None


def load_model(path: str | Path, issues: list[Issue]) -> list[Frame]:
    """Frames with their joint coordinates (metres). Problems go to `issues`."""
    from openpyxl import load_workbook

    path = Path(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # unreadable / not an xlsx
        issues.append(Issue("M03", f"Connectivity file cannot be read: {exc}", path.name))
        return []

    try:
        if FRAME_SHEET not in wb.sheetnames or JOINT_SHEET not in wb.sheetnames:
            issues.append(
                Issue("M03", f"Sheets '{FRAME_SHEET}' and '{JOINT_SHEET}' are required.", path.name)
            )
            return []
        frame_rows = list(wb[FRAME_SHEET].iter_rows(values_only=True))
        joint_rows = list(wb[JOINT_SHEET].iter_rows(values_only=True))
    finally:
        wb.close()

    fh = _find_header(frame_rows, ["Frame", "JointI", "JointJ", "Length"])
    jh = _find_header(joint_rows, ["Joint", "GlobalX", "GlobalY", "GlobalZ"])
    if fh is None or jh is None:
        issues.append(Issue("M03", "Header columns of the Excel sheets were not found.", path.name))
        return []

    f_row, f_col = fh
    units = frame_rows[f_row + 1] if f_row + 1 < len(frame_rows) else ()
    unit = str(units[f_col["Length"]]).strip() if len(units) > f_col["Length"] else ""
    if unit != "m":
        issues.append(Issue("M04", f"Length unit must be 'm' (found '{unit}').", path.name))
        return []

    j_row, j_col = jh
    joints: dict[str, tuple[float, float, float]] = {}
    for row in joint_rows[j_row + 2 :]:
        if row[j_col["Joint"]] is None:
            continue
        joints[_name(row[j_col["Joint"]])] = (
            float(row[j_col["GlobalX"]]),
            float(row[j_col["GlobalY"]]),
            float(row[j_col["GlobalZ"]]),
        )

    frames: list[Frame] = []
    for row in frame_rows[f_row + 2 :]:
        if row[f_col["Frame"]] is None:
            continue
        name = _name(row[f_col["Frame"]])
        ji, jj = _name(row[f_col["JointI"]]), _name(row[f_col["JointJ"]])
        missing = [j for j in (ji, jj) if j not in joints]
        if missing:
            issues.append(
                Issue("M05", f"Frame {name}: joint {', '.join(missing)} is not in the Joint sheet.", path.name)
            )
            continue
        frames.append(Frame(name, joints[ji], joints[jj], float(row[f_col["Length"]])))
    return frames


# ----------------------------------------------------------------------
# Cache: the Excel file is read once and reused while the file is unchanged
# (changing z0 / tolerance / coefficients must not read it again).
# ----------------------------------------------------------------------

_cache_lock = threading.Lock()
_cache: dict[tuple[str, int, int], tuple[list[Frame], list[Issue]]] = {}
_CACHE_MAX = 2


def load_model_cached(path: str | Path, issues: list[Issue]) -> list[Frame]:
    """Same as `load_model`, but keyed on (path, modification time, size).

    Safe to call from several threads: the lock is held while the file is read,
    so a second caller waits for the first read instead of repeating it.
    """
    path = Path(path)
    try:
        stat = path.stat()
    except OSError:
        return load_model(path, issues)  # reports the problem as M03

    key = (str(path.resolve()), stat.st_mtime_ns, stat.st_size)
    with _cache_lock:
        hit = _cache.get(key)
        if hit is None:
            local: list[Issue] = []
            frames = load_model(path, local)
            hit = (frames, local)
            if len(_cache) >= _CACHE_MAX:
                _cache.pop(next(iter(_cache)))
            _cache[key] = hit

    issues.extend(hit[1])
    return list(hit[0])
