"""Write SAP2000 import tables using the original MATLAB Excel templates."""
from __future__ import annotations

from pathlib import Path

from .pipeline import ConversionResult

JOINT_FILE = "SAP2000 Joint Coordinates.xlsx"
LOAD_FILE = "SAP2000 Distributed Loads.xlsx"
TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"


def _frame_value(name: str):
    """Frame name for the Excel cell.

    Written as a number only when the name is the canonical form of an integer
    ("9", "128"), exactly as the MATLAB tool wrote it. Every other name ("007", "B12",
    "9.5") stays text, so the name SAP receives is always the name it exported.
    """
    if name.isascii() and name.isdigit() and str(int(name)) == name:
        return int(name)
    return name


def write_outputs(result: ConversionResult, out_dir: str | Path) -> list[Path]:
    """Write both tables while preserving the embedded MATLAB workbook formatting."""
    from openpyxl import load_workbook

    if not result.ok:
        raise ValueError("The conversion has errors; no output is written.")

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    joint_path = out / JOINT_FILE
    wb = load_workbook(TEMPLATE_DIR / JOINT_FILE)
    ws = wb["Joint Coordinates"]
    for row_index, joint in enumerate(result.joints, start=4):
        values = [joint.name, "GLOBAL", "Cartesian", joint.x, joint.y, None,
                  joint.z, "Yes", None]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=column, value=value)
    wb.save(joint_path)
    wb.close()

    load_path = out / LOAD_FILE
    wb = load_workbook(TEMPLATE_DIR / LOAD_FILE)
    ws = wb["Frame Loads - Distributed"]
    for row_index, record in enumerate(result.loads, start=4):
        values = [_frame_value(record.frame), record.load_pat, "GLOBAL", "Force",
                  record.direction, "RelDist", record.rel_a, record.rel_b,
                  record.abs_a, record.abs_b, record.value, record.value,
                  None, None, None]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=column, value=value)
    wb.save(load_path)
    wb.close()

    return [joint_path, load_path]
