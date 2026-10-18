"""Excel import / export of the allowed-sections table (Auto Design, step 2).

Pure logic, no Qt. `openpyxl` is imported lazily, so the application
starts (and the rest of the wizard works) even if it is missing.

File layout: one sheet, header row `HEADERS`, one row per section,
Yes / No in the four "Is Allowed for ..." columns.
"""
from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from src.plugins.auto_design.params import SectionPermission

HEADERS = (
    "Section",
    "Is Allowed for Column",
    "Is Allowed for Beam",
    "Is Allowed for Horizontal Brace",
    "Is Allowed for Vertical Brace",
)

_YES = {"yes", "y", "true", "1", "x"}
_NO = {"no", "n", "false", "0", ""}


class SectionsExcelError(Exception):
    """A problem the user can understand (message is shown as is)."""


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise SectionsExcelError(
            "The 'openpyxl' package is required for Excel import and export."
        ) from exc

    return openpyxl


def parse_yes_no(value) -> bool | None:
    """Yes / No from a cell value; empty means No; unknown text gives None."""

    if value is None:
        return False

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        return bool(value) if value in (0, 1) else None

    text = str(value).strip().lower()

    if text in _YES:
        return True

    if text in _NO:
        return False

    return None


def _yes_no(flag: bool) -> str:
    return "Yes" if flag else "No"


# ---------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------


def export_sections(path: str | Path, sections: Iterable[SectionPermission]) -> None:
    """Write `sections` to an .xlsx file."""

    openpyxl = _openpyxl()

    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = list(sections)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sections"
    sheet.append(list(HEADERS))

    for section in rows:
        sheet.append(
            [
                section.name,
                _yes_no(section.allowedForColumn),
                _yes_no(section.allowedForBeam),
                _yes_no(section.allowedForHorizontalBrace),
                _yes_no(section.allowedForVerticalBrace),
            ]
        )

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    widest = max((len(s.name) for s in rows), default=10)
    sheet.column_dimensions["A"].width = max(18, widest + 2)

    for letter in ("B", "C", "D", "E"):
        sheet.column_dimensions[letter].width = 32

        for cell in sheet[letter][1:]:
            cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"

    if rows:
        # Drop-down in Excel so the user cannot type anything but Yes / No.
        validation = DataValidation(
            type="list", formula1='"Yes,No"', allow_blank=True
        )
        sheet.add_data_validation(validation)
        validation.add(f"B2:E{len(rows) + 1}")

    try:
        workbook.save(str(path))
    except OSError as exc:
        raise SectionsExcelError(
            f"Could not write the file (is it open in Excel?): {exc}"
        ) from exc


# ---------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------


def _cell(row, index: int):
    return row[index] if index < len(row) else None


def _find_columns(header) -> list[int]:
    if not header:
        raise SectionsExcelError("The first sheet of the file is empty.")

    names = {
        str(value).strip().lower(): index
        for index, value in enumerate(header)
        if value is not None
    }
    missing = [title for title in HEADERS if title.lower() not in names]

    if missing:
        raise SectionsExcelError(
            "Missing column(s) in the first row: " + ", ".join(missing) + "."
        )

    return [names[title.lower()] for title in HEADERS]


def _section_name(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)

    return str(value).strip()


def import_sections(path: str | Path) -> dict[str, tuple[bool, ...]]:
    """Read an .xlsx file: section name -> flags in the order of `HEADERS[1:]`
    (column, beam, horizontal brace, vertical brace).

    Nothing is returned when any Yes / No cell is unrecognised, so a typo
    never half-applies.
    """

    openpyxl = _openpyxl()

    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise SectionsExcelError(f"Could not open the Excel file: {exc}") from exc

    result: dict[str, tuple[bool, ...]] = {}
    problems: list[str] = []

    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        columns = _find_columns(next(rows, None))

        for number, row in enumerate(rows, start=2):
            name = _cell(row, columns[0])

            if name is None or not str(name).strip():
                continue

            flags = []

            for index, title in zip(columns[1:], HEADERS[1:]):
                value = _cell(row, index)
                flag = parse_yes_no(value)

                if flag is None:
                    problems.append(f"row {number}, {title}: {value!r}")
                    flag = False

                flags.append(flag)

            result[_section_name(name)] = tuple(flags)
    except SectionsExcelError:
        raise
    except Exception as exc:
        raise SectionsExcelError(f"Could not read the Excel file: {exc}") from exc
    finally:
        workbook.close()

    if problems:
        shown = "; ".join(problems[:5])
        more = f" (and {len(problems) - 5} more)" if len(problems) > 5 else ""
        raise SectionsExcelError(f"Use Yes or No only. Problem at {shown}{more}.")

    if not result:
        raise SectionsExcelError("The file contains no sections.")

    return result