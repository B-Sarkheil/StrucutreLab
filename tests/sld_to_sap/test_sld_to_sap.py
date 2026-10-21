"""Tests of the SLD to SAP services, using the worked examples of the docs (section 10)."""
from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from src.plugins.sld_to_sap.params import SldParams
from src.plugins.sld_to_sap.services.exporter import write_outputs
from src.plugins.sld_to_sap.services.geometry import round_mm
from src.plugins.sld_to_sap.services.pipeline import convert

HEADER = "Index,X1,Y1,X2,Y2,X3,Y3,X4,Y4,Text1,Text2,Text3"
ORIGIN_Z = "1,1823237.249,1685189.731,,,,,,,Z,2441.3,"
PARAMS = SldParams(z0=2433.7, tolerance_mm=10, eqz_coef=0.2)


def make_model(path, joints, frames):
    wb = Workbook()
    ws = wb.active
    ws.title = "Connectivity - Frame"
    ws.append(["TABLE:  Connectivity - Frame"])
    ws.append(["Frame", "JointI", "JointJ", "IsCurved", "Length"])
    ws.append(["Text", "Text", "Text", "Yes/No", "m"])
    for f in frames:
        ws.append(list(f))
    wj = wb.create_sheet("Joint Coordinates")
    wj.append(["TABLE:  Joint Coordinates"])
    wj.append(["Joint", "CoordSys", "CoordType", "XorR", "Y", "Z", "GlobalX", "GlobalY", "GlobalZ"])
    wj.append(["Text"] * 9)
    for name, x, y, z in joints:
        wj.append([name, "GLOBAL", "Cartesian", x, y, z, x, y, z])
    wb.save(path)


@pytest.fixture
def model(tmp_path):
    p = tmp_path / "CONNECTIVITY.xlsx"
    make_model(
        p,
        joints=[("73", 52.1, 0, 2441.3), ("74", 52.1, 5, 2441.3), ("80", 52.1, 10, 2441.3)],
        frames=[("66", "73", "74", "No", 5.0), ("67", "80", "74", "No", 5.0)],
    )
    return p


def write_csv(tmp_path, name, *lines):
    f = tmp_path / name
    f.write_text("\n".join([HEADER, *lines]) + "\n", encoding="utf-8")
    return f


def test_round_mm_away_from_zero():
    assert round_mm(0.0005) == 0.001
    assert round_mm(-0.0005) == -0.001
    assert round_mm(2.5) == 2.5


def test_plan_load_and_point(tmp_path, model):
    f = write_csv(
        tmp_path,
        "Z2441.300.csv",
        ORIGIN_Z,
        '2,1873337.249,1687954.731,,,,,,,"P015","","",',
        "8,1875253.131,1685889.731,1875421.367,1685889.731,1875421.367,1686889.731,"
        "1875253.131,1686889.731,V=1000,LS=500,",
    )
    r = convert(PARAMS, [f], model)
    assert r.ok, [str(i) for i in r.issues]

    j = r.joints[0]
    assert (j.name, j.x, j.y, j.z) == ("P015", 50.1, 2.765, 2441.3)

    assert [(x.load_pat, x.direction) for x in r.loads] == [
        ("PIPE.OPER1", "Gravity"), ("PIPE.OPER2", "Gravity"), ("EQ.Z", "Gravity"),
        ("FRX", "X"), ("FRY", "Y"), ("ANC.OCCY", "Y"),
    ]
    first = r.loads[0]
    assert first.frame == "66"
    assert first.abs_a == pytest.approx(0.7) and first.abs_b == pytest.approx(1.7)
    assert first.rel_a == pytest.approx(0.14) and first.rel_b == pytest.approx(0.34)
    assert r.loads[2].value == pytest.approx(0.2)
    assert r.loads[3].value == pytest.approx(0.15)


def test_reversed_frame_distances(tmp_path, model):
    # frame 67 runs from y=10 to y=5 (I > J); load on y = 6..7
    f = write_csv(
        tmp_path,
        "Z.csv",
        ORIGIN_Z,
        "2,1875253.131,1691189.731,1875421.367,1691189.731,1875421.367,1692189.731,"
        "1875253.131,1692189.731,V=1000,",
    )
    r = convert(PARAMS, [f], model)
    assert r.ok, [str(i) for i in r.issues]
    assert r.loads[0].frame == "67"
    assert r.loads[0].abs_a == pytest.approx(3.0) and r.loads[0].abs_b == pytest.approx(4.0)


def test_load_sticking_out_is_rejected_for_reversed_frame(tmp_path, model):
    # y = 4..6 overlaps frame 67 (5..10) only partly: MATLAB would accept it
    f = write_csv(
        tmp_path,
        "Z.csv",
        ORIGIN_Z,
        "2,1875253.131,1689189.731,1875421.367,1689189.731,1875421.367,1691189.731,"
        "1875253.131,1691189.731,V=1000,",
    )
    r = convert(PARAMS, [f], model)
    assert not r.ok
    assert {i.code for i in r.errors} == {"M01"}


def test_structural_errors_are_all_collected(tmp_path, model):
    f = write_csv(
        tmp_path,
        "Z.csv",
        ORIGIN_Z,
        '2,1873337.249,1687954.731,,,,,,,"P1","","",',
        '3,1873337.249,1687954.731,,,,,,,"P1","","",',  # repeated name
        "4,1875253.131,1685889.731,1875421.367,1685889.731,1875421.367,1686889.731,"
        "1875253.131,1686889.731,,,",  # no text
        "5,1875253.131,1685889.731,1875421.367,1685889.731,1875421.367,1686889.731,"
        "1875253.131,1686889.731,V=1500kg,1500,XX=3",
    )
    r = convert(PARAMS, [f], model)
    codes = sorted(i.code for i in r.errors)
    assert codes == ["S05", "S10", "S13", "S14", "S15"]


def test_missing_origin(tmp_path, model):
    f = write_csv(tmp_path, "A.csv", '2,1,1,,,,,,,"P1","",')
    r = convert(PARAMS, [f], model)
    assert [i.code for i in r.errors] == ["S01"]


def test_export_roundtrip(tmp_path, model):
    f = write_csv(
        tmp_path,
        "Z.csv",
        ORIGIN_Z,
        '2,1873337.249,1687954.731,,,,,,,"P015","",',
        "8,1875253.131,1685889.731,1875421.367,1685889.731,1875421.367,1686889.731,"
        "1875253.131,1686889.731,V=1000,",
    )
    r = convert(PARAMS, [f], model)
    paths = write_outputs(r, tmp_path / "out")
    wb = load_workbook(paths[1])
    ws = wb["Frame Loads - Distributed"]
    assert ws["A4"].value == 66 and ws["B4"].value == "PIPE.OPER1"
    assert ws.max_row == 3 + 5
    assert load_workbook(paths[0])["Joint Coordinates"]["A4"].value == "P015"

    bad = convert(PARAMS, [write_csv(tmp_path, "B.csv", '2,1,1,,,,,,,"P1","",')], model)
    with pytest.raises(ValueError):
        write_outputs(bad, tmp_path / "out2")


def test_check_structure_needs_no_model(tmp_path):
    from src.plugins.sld_to_sap.services.pipeline import check_structure

    a = write_csv(tmp_path, "A.csv", ORIGIN_Z, '2,1,1,,,,,,,"P1","",', '3,1,1,,,,,,,"P1","",')
    b = write_csv(tmp_path, "B.csv", '2,1,1,,,,,,,"P2","",')
    r = check_structure([a, b])
    assert sorted((i.file, i.code) for i in r.errors) == [("A.csv", "S10"), ("B.csv", "S01")]
